"""Tests for src/excel_report.py: workbook structure, the security guard, and
correctness of the Pricing Simulator's live Excel formulas.

No LibreOffice/Excel engine is available in this environment to actually
recalculate a .xlsx, so formula correctness is verified two ways (per
ARCHITECTURE.md's "reimplement formula logic in Python" fallback):

1. `compute_expected_simulator_outputs` (production code) recomputes every
   simulator output straight from the same baseline data.
2. `_MiniFormulaEvaluator` below is a small, independent, from-scratch
   Excel-formula interpreter (cell refs, ranges, SUM/MAX/IF/SUMPRODUCT,
   +-*/, comparisons) that evaluates the *actual formula strings* written
   into the workbook, resolving cell-to-cell dependencies recursively. It
   shares no code with `excel_report.py`'s formula-building logic, so
   agreement between the two is a real cross-check, not a tautology.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest
from openpyxl.utils import column_index_from_string, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

import excel_report as er  # noqa: E402
from db import DB_PATH, get_conn  # noqa: E402

requires_db = pytest.mark.skipif(
    not DB_PATH.exists(), reason="data/agentops.db not generated yet — run `make data` first"
)

EXPECTED_SHEETS = [
    "README", "KPI Dashboard", "Monthly KPIs", "Plan Economics", "Model x Task",
    "Customers", "Pricing Simulator", "Pivot-ready Data", "_check",
]


@pytest.fixture(scope="module")
def workbook_path(tmp_path_factory):
    out = tmp_path_factory.mktemp("excel_out") / "AgentOps_Report.xlsx"
    er.build_workbook(out)
    return out


@pytest.fixture(scope="module")
def workbook(workbook_path):
    return openpyxl.load_workbook(workbook_path)


# --------------------------------------------------------------------- sanitize_cell


@pytest.mark.parametrize(
    "raw",
    [
        "=cmd|' /C calc'!A0",
        "=1+1",
        "+1+1",
        "-2+3",
        "@SUM(1,2)",
        "\tsneaky",
        "\rsneaky",
    ],
)
def test_sanitize_cell_neutralises_dangerous_prefixes(raw):
    out = er.sanitize_cell(raw)
    assert isinstance(out, str)
    assert out.startswith("'")
    assert out == "'" + raw


@pytest.mark.parametrize(
    "raw",
    ["Acme Corp", "North America", "Mid-Market", "test_generation", "a-b-c not a formula"],
)
def test_sanitize_cell_leaves_normal_strings_untouched(raw):
    assert er.sanitize_cell(raw) == raw


@pytest.mark.parametrize("raw", [-5, -5.25, 0, 42, 3.14])
def test_sanitize_cell_leaves_numbers_untouched(raw):
    out = er.sanitize_cell(raw)
    assert out == raw
    assert type(out) is type(raw)


def test_sanitize_cell_passes_through_none():
    assert er.sanitize_cell(None) is None


# --------------------------------------------------------------------- workbook structure


@requires_db
def test_expected_sheets_present(workbook):
    assert workbook.sheetnames == EXPECTED_SHEETS


@requires_db
def test_check_sheet_is_hidden(workbook):
    assert workbook["_check"].sheet_state == "hidden"


@requires_db
def test_readme_mentions_synthetic_data(workbook):
    ws = workbook["README"]
    text = " ".join(str(c.value) for row in ws.iter_rows() for c in row if c.value)
    assert "SYNTHETIC" in text.upper()


@requires_db
@pytest.mark.parametrize("sheet,table_name", [
    ("Monthly KPIs", "MonthlyKPIs"),
    ("Plan Economics", "PlanEconomics"),
    ("Model x Task", "ModelTask"),
    ("Customers", "Customers"),
    ("Pivot-ready Data", "PivotData"),
])
def test_expected_excel_tables_exist(workbook, sheet, table_name):
    ws = workbook[sheet]
    assert table_name in ws.tables
    assert ws.freeze_panes == "A2"


@requires_db
def test_customers_sheet_has_conditional_formatting(workbook):
    ws = workbook["Customers"]
    assert len(ws.conditional_formatting) >= 2


@requires_db
def test_kpi_dashboard_has_charts(workbook):
    ws = workbook["KPI Dashboard"]
    assert len(ws._charts) >= 2


# --------------------------------------------------------------------- Pricing Simulator: formulas not values


@requires_db
def test_simulator_output_cells_are_formulas(workbook):
    ws = workbook["Pricing Simulator"]
    formula_cells = [
        c.coordinate
        for row in ws.iter_rows()
        for c in row
        if isinstance(c.value, str) and c.value.startswith("=")
    ]
    # The showpiece must have a non-trivial number of live formulas (per-plan
    # outputs, totals, and the 4-column scenario grid).
    assert len(formula_cells) >= 25


@requires_db
def test_simulator_input_cells_are_plain_values_not_formulas(workbook):
    ws = workbook["Pricing Simulator"]
    # Plan price inputs are yellow INPUT_FILL cells holding literal numbers.
    for row in ws.iter_rows():
        for c in row:
            if c.fill and c.fill.fgColor and c.fill.fgColor.rgb == "00FFF2CC":
                assert not (isinstance(c.value, str) and c.value.startswith("="))
                assert isinstance(c.value, (int, float))


@requires_db
def test_simulator_has_data_validation(workbook):
    ws = workbook["Pricing Simulator"]
    assert len(ws.data_validations.dataValidation) > 0


# --------------------------------------------------------------------- Pricing Simulator: formula correctness


CELL_RE = re.compile(r"\$?[A-Z]{1,2}\$?\d+")


class _MiniFormulaEvaluator:
    """Tiny, independent Excel-formula interpreter for the subset of syntax
    used in the Pricing Simulator sheet (cell refs, ranges, + - * /,
    comparisons, and SUM/MAX/IF/SUMPRODUCT). Built from scratch rather than
    reusing excel_report.py's formula-generation code, so agreement with
    `compute_expected_simulator_outputs` is a genuine independent check.
    """

    def __init__(self, ws):
        self.ws = ws
        self.cache: dict[str, object] = {}

    def get_cell(self, coord: str):
        coord = coord.replace("$", "")
        if coord in self.cache:
            return self.cache[coord]
        val = self.ws[coord].value
        if isinstance(val, str) and val.startswith("="):
            result = _Parser(val[1:], self).parse()
        else:
            result = val if val is not None else 0
        self.cache[coord] = result
        return result

    @staticmethod
    def expand_range(a: str, b: str) -> list[str]:
        col1 = column_index_from_string(re.match(r"[A-Z]+", a).group())
        row1 = int(re.search(r"\d+", a).group())
        col2 = column_index_from_string(re.match(r"[A-Z]+", b).group())
        row2 = int(re.search(r"\d+", b).group())
        return [
            f"{get_column_letter(c)}{r}"
            for r in range(row1, row2 + 1)
            for c in range(col1, col2 + 1)
        ]


class _Parser:
    def __init__(self, s: str, evaluator: _MiniFormulaEvaluator):
        self.s = s
        self.pos = 0
        self.ev = evaluator

    def next_tok(self):
        while self.pos < len(self.s) and self.s[self.pos] == " ":
            self.pos += 1
        if self.pos >= len(self.s):
            return None
        m = re.match(r"\$?[A-Z]{1,2}\$?\d+", self.s[self.pos:])
        if m:
            self.pos += m.end()
            return ("cell", m.group().replace("$", ""))
        m = re.match(r"[A-Za-z_]+", self.s[self.pos:])
        if m:
            self.pos += m.end()
            return ("name", m.group())
        m = re.match(r"\d+\.?\d*", self.s[self.pos:])
        if m:
            self.pos += m.end()
            return ("num", float(m.group()))
        ch = self.s[self.pos]
        self.pos += 1
        if ch in ("<", ">", "=") and self.pos < len(self.s) and self.s[self.pos] in ("=", ">"):
            ch2 = self.s[self.pos]
            self.pos += 1
            return ("op", ch + ch2)
        return ("op", ch)

    def parse(self):
        return self.parse_comparison()

    def parse_comparison(self):
        left = self.parse_addsub()
        save = self.pos
        t = self.next_tok()
        if t and t[0] == "op" and t[1] in ("=", "<>", ">=", "<=", ">", "<"):
            right = self.parse_addsub()
            return {
                "=": left == right, "<>": left != right, ">=": left >= right,
                "<=": left <= right, ">": left > right, "<": left < right,
            }[t[1]]
        self.pos = save
        return left

    def parse_addsub(self):
        val = self.parse_muldiv()
        while True:
            save = self.pos
            t = self.next_tok()
            if t and t[0] == "op" and t[1] in ("+", "-"):
                rhs = self.parse_muldiv()
                val = val + rhs if t[1] == "+" else val - rhs
            else:
                self.pos = save
                break
        return val

    def parse_muldiv(self):
        val = self.parse_unary()
        while True:
            save = self.pos
            t = self.next_tok()
            if t and t[0] == "op" and t[1] in ("*", "/"):
                rhs = self.parse_unary()
                val = val * rhs if t[1] == "*" else val / rhs
            else:
                self.pos = save
                break
        return val

    def parse_unary(self):
        save = self.pos
        t = self.next_tok()
        if t and t[0] == "op" and t[1] == "-":
            return -self.parse_unary()
        self.pos = save
        return self.parse_atom()

    def parse_atom(self):
        t = self.next_tok()
        if t[0] == "num":
            return t[1]
        if t[0] == "cell":
            save = self.pos
            nt = self.next_tok()
            if nt and nt == ("op", ":"):
                nt2 = self.next_tok()
                return ("RANGE", t[1], nt2[1])
            self.pos = save
            return self.ev.get_cell(t[1])
        if t[0] == "name":
            fname = t[1]
            assert self.next_tok() == ("op", "(")
            args = self.parse_args()
            assert self.next_tok() == ("op", ")")
            return self.call(fname, args)
        if t[0] == "op" and t[1] == "(":
            val = self.parse_comparison()
            assert self.next_tok() == ("op", ")")
            return val
        raise ValueError(f"unexpected token {t} in {self.s!r}")

    def parse_args(self):
        args = [self.parse_arg()]
        while True:
            save = self.pos
            t = self.next_tok()
            if t == ("op", ","):
                args.append(self.parse_arg())
            else:
                self.pos = save
                break
        return args

    def parse_arg(self):
        save = self.pos
        t = self.next_tok()
        if t and t[0] == "cell":
            save2 = self.pos
            nt = self.next_tok()
            if nt == ("op", ":"):
                nt2 = self.next_tok()
                return ("RANGE", t[1], nt2[1])
            self.pos = save2
        self.pos = save
        return self.parse_comparison()

    def call(self, fname, args):
        if fname == "SUM":
            _, a, b = args[0]
            return sum(self.ev.get_cell(c) for c in self.ev.expand_range(a, b))
        if fname == "MAX":
            return max(args)
        if fname == "IF":
            cond, t, f = args
            return t if cond else f
        if fname == "SUMPRODUCT":
            ranges = [self.ev.expand_range(a[1], a[2]) for a in args]
            columns = [[self.ev.get_cell(c) for c in rng] for rng in ranges]
            return sum(v1 * v2 for v1, v2 in zip(*columns))
        raise ValueError(f"unsupported function {fname}")


def _find_header_row(ws, needle: str) -> int:
    """Locate a section header by its text (substring match), not a hardcoded
    row number, so tests keep working if the sheet grows/shrinks above it."""
    for row in ws.iter_rows(min_col=1, max_col=1):
        for cell in row:
            if isinstance(cell.value, str) and needle in cell.value:
                return cell.row
    raise AssertionError(f"header {needle!r} not found in column A")


def _plan_rows_below(ws, header_row: int) -> dict[str, int]:
    """The 3 plan-label rows (Pro/Team/Enterprise) immediately below a section header."""
    rows: dict[str, int] = {}
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + 8, min_col=1, max_col=1):
        for cell in row:
            if cell.value in er.PAID_PLANS:
                rows[cell.value] = cell.row
    return rows


@requires_db
def test_simulator_formulas_evaluate_to_expected_default_state(workbook_path):
    """Evaluate the actual formula strings in the workbook (default input
    state) with an independent interpreter and compare against
    `compute_expected_simulator_outputs`, which recomputes the same figures
    straight from the underlying data — a real cross-check of the formula
    wiring, not a tautology. Also confirms Δ margin (projected vs the locked
    baseline block) is exactly 0 for every plan and blended at default inputs."""
    with get_conn() as conn:
        baseline = er.fetch_simulator_baseline(conn)
    expected = er.compute_expected_simulator_outputs(baseline)

    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Pricing Simulator"]
    evaluator = _MiniFormulaEvaluator(ws)

    header_row = _find_header_row(ws, "Projected outputs")
    plan_rows = _plan_rows_below(ws, header_row)
    assert set(plan_rows) == set(er.PAID_PLANS)

    for plan, row in plan_rows.items():
        exp = expected["per_plan"][plan]
        revenue = evaluator.get_cell(f"D{row}")
        cost = evaluator.get_cell(f"E{row}")
        margin = evaluator.get_cell(f"F{row}")
        delta_pp = evaluator.get_cell(f"G{row}")
        assert revenue == pytest.approx(exp["revenue"], rel=1e-3)
        assert cost == pytest.approx(exp["cost"], rel=1e-3)
        assert margin == pytest.approx(exp["margin"], abs=1e-3)
        # Orchestrator-mandated invariant: at default inputs, the projected
        # block and the locked baseline block use identical formula
        # structure over identical values, so delta must be exactly 0.
        assert delta_pp == pytest.approx(0.0, abs=1e-9)

    total_row = max(plan_rows.values()) + 1
    total_revenue = evaluator.get_cell(f"D{total_row}")
    total_cost = evaluator.get_cell(f"E{total_row}")
    total_margin = evaluator.get_cell(f"F{total_row}")
    total_delta_pp = evaluator.get_cell(f"G{total_row}")
    assert total_revenue == pytest.approx(expected["total_revenue"], rel=1e-3)
    assert total_cost == pytest.approx(expected["total_cost"], rel=1e-3)
    assert total_margin == pytest.approx(expected["total_margin"], abs=1e-3)
    assert total_delta_pp == pytest.approx(0.0, abs=1e-9)


@requires_db
def test_default_pct_routed_is_zero(workbook):
    """Status quo has no model routing — the '% routed' input must default to 0."""
    ws = workbook["Pricing Simulator"]
    header_row = _find_header_row(ws, "% of simple-task frontier runs routed")
    assert ws.cell(row=header_row, column=2).value == 0


@requires_db
def test_baseline_locked_block_present_and_gray(workbook):
    ws = workbook["Pricing Simulator"]
    header_row = _find_header_row(ws, "Baseline outputs (locked model")
    plan_rows = _plan_rows_below(ws, header_row)
    assert set(plan_rows) == set(er.PAID_PLANS)
    # Baseline output cells are locked (not the yellow INPUT_FILL).
    for row in plan_rows.values():
        fill = ws.cell(row=row, column=4).fill  # revenue cell
        assert fill.fgColor.rgb != "00FFF2CC"


@requires_db
def test_actual_margin_reference_and_jensen_note_present(workbook):
    ws = workbook["Pricing Simulator"]
    title_row = _find_header_row(ws, "Projected outputs")
    header_row = title_row + 1  # column headers sit one row below the section title
    # "Actual margin in data (reference)" column header.
    header_texts = [ws.cell(row=header_row, column=c).value for c in range(1, 9)]
    assert any(isinstance(h, str) and "Actual margin in data" in h for h in header_texts)
    plan_rows = _plan_rows_below(ws, header_row)
    for row in plan_rows.values():
        val = ws.cell(row=row, column=8).value
        assert val is None or isinstance(val, (int, float))
    all_text = " ".join(
        str(c.value) for r in ws.iter_rows() for c in r if isinstance(c.value, str)
    )
    assert "Jensen" in all_text


@requires_db
def test_price_elasticity_input_present_with_ab_test_default(workbook):
    ws = workbook["Pricing Simulator"]
    header_row = _find_header_row(ws, "Plan pricing inputs")
    plan_rows = _plan_rows_below(ws, header_row)
    with get_conn() as conn:
        baseline = er.fetch_simulator_baseline(conn)
    expected_elasticity = float(baseline["elasticity_pro"])
    assert expected_elasticity < 0  # a price hike should reduce conversion/customers
    for row in plan_rows.values():
        elasticity_cell = ws.cell(row=row, column=5)
        assert elasticity_cell.value == pytest.approx(round(expected_elasticity, 4), abs=1e-4)


@requires_db
def test_raising_price_with_negative_elasticity_reduces_customers(workbook_path, tmp_path):
    """A live, end-to-end check that the elasticity wiring actually responds:
    bump Pro's price input, evaluate the *actual* formula cells, and confirm
    adjusted customers drop (elasticity is negative) while the baseline
    (locked) customer count is untouched."""
    wb = openpyxl.load_workbook(workbook_path)
    ws = wb["Pricing Simulator"]

    pricing_header = _find_header_row(ws, "Plan pricing inputs")
    plan_rows = _plan_rows_below(ws, pricing_header)
    pro_row = plan_rows["Pro"]
    ws.cell(row=pro_row, column=2, value=39.0)  # bump Pro price 29 -> 39

    bumped_path = tmp_path / "bumped.xlsx"
    wb.save(bumped_path)

    wb2 = openpyxl.load_workbook(bumped_path)
    ws2 = wb2["Pricing Simulator"]
    evaluator = _MiniFormulaEvaluator(ws2)

    projected_header = _find_header_row(ws2, "Projected outputs")
    baseline_header = _find_header_row(ws2, "Baseline outputs (locked model")
    projected_rows = _plan_rows_below(ws2, projected_header)
    baseline_rows = _plan_rows_below(ws2, baseline_header)

    projected_customers = evaluator.get_cell(f"B{projected_rows['Pro']}")
    baseline_customers = evaluator.get_cell(f"B{baseline_rows['Pro']}")

    assert projected_customers < baseline_customers
    assert baseline_customers == pytest.approx(397.0, abs=0.5)


@requires_db
def test_check_sheet_values_match_python_recompute(workbook):
    """The hidden `_check` sheet's TOTAL row should match
    `compute_expected_simulator_outputs` (same function that produced it —
    this guards against the sheet drifting from the function on a future edit)."""
    with get_conn() as conn:
        baseline = er.fetch_simulator_baseline(conn)
    expected = er.compute_expected_simulator_outputs(baseline)

    ws = workbook["_check"]
    values = {row[0].value: row for row in ws.iter_rows(min_row=4, max_col=8) if row[0].value}
    total_row = values["TOTAL"]
    assert total_row[3].value == pytest.approx(expected["total_revenue"], rel=1e-6)
    assert total_row[4].value == pytest.approx(expected["total_cost"], rel=1e-6)
    assert total_row[5].value == pytest.approx(expected["total_margin"], rel=1e-6)
    assert total_row[6].value == pytest.approx(expected["baseline_total_margin"], rel=1e-6)
    assert total_row[7].value == pytest.approx(0.0, abs=1e-6)


@requires_db
def test_scenario_comparison_block_present(workbook):
    ws = workbook["Pricing Simulator"]
    labels = [c.value for row in ws.iter_rows(min_col=1, max_col=1) for c in row]
    assert any(isinstance(v, str) and "Scenario comparison" in v for v in labels)
    header_cells = [c for row in ws.iter_rows() for c in row if c.value in ("Status Quo",)]
    assert header_cells, "expected a 'Status Quo' scenario column header"


# --------------------------------------------------------------------- Customers sheet activation logic


@requires_db
def test_customers_activation_matches_shared_definition():
    """Activation = success share of a customer's first 5 runs >= 3/5, per
    ARCHITECTURE.md's shared metric definitions."""
    with get_conn() as conn:
        df = er.fetch_customers(conn)
    activated_mask = df["first5_success_rate"] >= 0.6
    assert (df["activated"] == activated_mask).all()
